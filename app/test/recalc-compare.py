"""Compara els valors de les fórmules que ha posat el generador amb els que recalcula LibreOffice (recàlcul complet).
   ús: python3 recalc-compare.py <perfil-de-libreoffice> fitxer.xlsx [...]
   Les fórmules que depenen de TODAY() s'ometen (el rellotge dels tests és fix)."""
import sys, os, re, subprocess, shutil, math, datetime, tempfile
import openpyxl

profile, files = sys.argv[1], sys.argv[2:]

def near(a, b):
    if a in (None, '') and b in (None, ''): return True
    if isinstance(a, (int, float)) and isinstance(b, (int, float)): return math.isclose(a, b, rel_tol=1e-6, abs_tol=1e-6)
    return str(a if a is not None else '') == str(b if b is not None else '')

total = bad_total = 0
for f in files:
    out = tempfile.mkdtemp()
    subprocess.run(['soffice', f'-env:UserInstallation=file://{profile}', '--headless', '--convert-to', 'xlsx:Calc MS Excel 2007 XML', '--outdir', out, f], capture_output=True, timeout=240)
    lo = os.path.join(out, os.path.basename(f))
    wf = openpyxl.load_workbook(f)
    wc = openpyxl.load_workbook(f, data_only=True)
    wl = openpyxl.load_workbook(lo, data_only=True)
    n = bad = 0
    for ws in wf.worksheets:
        tainted = {c.coordinate for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith('=') and 'TODAY(' in c.value}
        for row in ws.iter_rows():
            for c in row:
                if not (isinstance(c.value, str) and c.value.startswith('=')): continue
                if 'TODAY(' in c.value: continue
                refs = set(re.findall(r'(?<![A-Za-z!])\$?([A-Z]{1,2})\$?(\d+)', c.value))
                if any(f'{a}{b}' in tainted for a, b in refs): continue
                n += 1
                mine, theirs = wc[ws.title][c.coordinate].value, wl[ws.title][c.coordinate].value
                if not near(mine, theirs):
                    bad += 1
                    if bad <= 10: print(f'  ✗ {ws.title}!{c.coordinate}: {c.value[:100]} generador={mine!r} libreoffice={theirs!r}')
    print(f'{os.path.basename(f)}: {n} fórmules, {bad} diferències')
    shutil.rmtree(out, ignore_errors=True)
    total += n; bad_total += bad
print(f'TOTAL {total} fórmules, {bad_total} diferències')
sys.exit(1 if bad_total else 0)
